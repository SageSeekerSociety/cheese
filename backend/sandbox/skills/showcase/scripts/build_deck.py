#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_deck.py — 把一份 JSON 幻灯片稿渲染成 16:9 的 .pptx。

用法:
    uv run --with python-pptx python build_deck.py spec.json out.pptx

设计沿用知是设计体系（references/design-system.md 的 --cx-* token）：
中性底、ink/text/muted 三级灰、单一琥珀强调、状态三色、同一套边距与栅格。
页面上的每块字都是 PPT 里可编辑的原生文本框（没有把文字画成图片），图表走
python-pptx 原生图表 API，表格是真的表格，讲者备注写进 notes。

中文字体：西文 Segoe UI、东亚 Microsoft YaHei。python-pptx 的 font.name 只写
<a:latin>，所以每处文字都另写 <a:ea>（必要时 <a:cs>），让 CJK 真的落到中文字体上，
而不是悄悄回退。

JSON 结构（字段都可缺省，缺了就不画）:
{
  "title": "...", "subtitle": "...", "meta": "...",
  "slides": [
    {"layout":"cover",     "eyebrow","title","subtitle",
                            "hero_from","hero_to","hero_cap","meta","notes"},
    {"layout":"section",   "eyebrow","title","subtitle","notes"},
    {"layout":"statement", "eyebrow","text","accent","note","notes"},
    {"layout":"bullets",   "eyebrow","title","items":[{"text","meta"}],"aside":{"label","value","note"},"notes"},
    {"layout":"two-col",   "eyebrow","title","cols":[{"title","tone","items":[...]}],"notes"},
    {"layout":"chart",     "eyebrow","title","takeaway","caption",
                            "chart":{"categories":[...],"series":[{"name","values":[...]}],"value_suffix","type"},"notes"},
    {"layout":"table",     "eyebrow","title","caption","table":{"columns":[...],"col_align":[...],"rows":[[...]]},"notes"},
    {"layout":"timeline",  "eyebrow","title","steps":[{"label","desc"}],"notes"},
    {"layout":"closing",   "eyebrow","title",
                            "lines":["...",{"text":"...","accent":true}],"meta","notes"}
  ]
}
封面 hero_from/hero_to 是「改前 → 改后」那对比：hero_from 划掉、hero_to 用琥珀，
是全页唯一的强调。closing 的某行写成 {"text","accent"} 就把琥珀落在那行上。
单元格里可以放字符串，也可以放 {"text":"已改","tone":"ok|warn|danger|neutral"} 画状态标记。
"""
import json
import math
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import (XL_CHART_TYPE, XL_LEGEND_POSITION,
                                 XL_LABEL_POSITION)
from pptx.oxml.ns import qn

# --------------------------------------------------------------------------
# 设计 token —— 与 references/design-system.md §2 一致
# --------------------------------------------------------------------------
C = {
    "canvas": "F7F8FA", "surface": "FFFFFF", "raised": "FFFFFF",
    "ink": "191A1C", "text": "36383C", "muted": "5A5E66", "faint": "747A82",
    "line": "ECEDEF", "line2": "E2E3E6", "fill": "F4F5F7", "fill2": "EEEEF1",
    "accent": "F57F17", "accent_press": "D96E0A", "accent_ink": "9A5413",
    "accent_wash": "FDF1E2", "on_accent": "FFFFFF",
    "ok": "1F9D55", "ok_ink": "12703A", "ok_wash": "E8F6EE",
    "warn": "E8901C", "warn_ink": "8F5406", "warn_wash": "FAF0DC",
    "danger": "DC2626", "danger_ink": "B91C1C", "danger_wash": "FDECEC",
    "chart1": "3D6BB3", "chart2": "2A9D8F", "chart3": "D9822B",
    "chart4": "8A63C2", "chart5": "C9506F", "chart6": "6B7B8C",
}
TONE_INK = {"ok": C["ok_ink"], "warn": C["warn_ink"],
            "danger": C["danger_ink"], "neutral": C["muted"]}
TONE_MARK = {"ok": C["ok"], "warn": C["warn"],
             "danger": C["danger"], "neutral": C["faint"]}
TONE_WASH = {"ok": C["ok_wash"], "warn": C["warn_wash"],
             "danger": C["danger_wash"], "note": C["fill"]}

# 字号阶梯（pt）—— 正文 ≥ 18pt，与页面版的 13/16/22/34 等价放大
F_EYEBROW, F_TITLE, F_BODY = 13, 27, 18
F_MUTED, F_CAPTION, F_FOOT = 15, 13, 11
LATIN, EA = "Segoe UI", "Microsoft YaHei"

# 16:9 栅格（英寸）
SW, SH = 13.333, 7.5
MX = 0.75                 # 左右边距
CW = SW - 2 * MX          # 内容宽 11.833
TOP = 0.62                # 内容起点
CHARTS = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED,
          "bar": XL_CHART_TYPE.BAR_CLUSTERED,
          "line": XL_CHART_TYPE.LINE_MARKERS}


def rgb(name_or_hex):
    h = C.get(name_or_hex, name_or_hex)
    return RGBColor.from_string(h)


# --------------------------------------------------------------------------
# 底层：字体、文本、形状
# --------------------------------------------------------------------------
def _set_face(parent, tag, name):
    el = parent.find(qn(tag))
    if el is None:
        el = parent.makeelement(qn(tag), {})
        parent.append(el)
    el.set("typeface", name)


def _style_run(run, size, color, bold=False, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = LATIN
    rPr = run._r.get_or_add_rPr()
    _set_face(rPr, "a:ea", EA)     # 关键：python-pptx 不会写 a:ea
    _set_face(rPr, "a:cs", LATIN)


def _run(p, text, size, color, bold=False):
    r = p.add_run()
    r.text = text
    _style_run(r, size, color, bold)
    return r


def para(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT,
         line=None, space_before=None, space_after=None, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    if line is not None:
        p.line_spacing = line
    if space_before is not None:
        p.space_before = Pt(space_before)
    if space_after is not None:
        p.space_after = Pt(space_after)
    if text:
        _run(p, text, size, color, bold)
    return p


def tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=0.75, radius=None):
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    # 主题的 <p:style> 带 effectRef（outerShdw），留下会盖过 shadow.inherit；
    # 形状的填充/描边都已显式设过，整块 style 去掉更干净。
    style = sp._element.find(qn("p:style"))
    if style is not None:
        sp._element.remove(style)
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def oval(slide, cx, cy, d, fill):
    return shape(slide, MSO_SHAPE.OVAL, cx - d / 2.0, cy - d / 2.0, d, d, fill=fill)


def hline(slide, x, y, w, color, line_w=1.0):
    ln = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(line_w)
    ln.shadow.inherit = False
    style = ln._element.find(qn("p:style"))
    if style is not None:
        ln._element.remove(style)
    return ln


# --------------------------------------------------------------------------
# 页面骨架
# --------------------------------------------------------------------------
def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    return prs


def add_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])   # 空白版式
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb("canvas")
    return s


def set_notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def footer(slide, idx, total, name):
    hline(slide, MX, 6.94, CW, rgb("line"), 0.75)
    tf = tb(slide, MX, 7.02, CW * 0.72, 0.3)
    para(tf, name, F_FOOT, rgb("faint"), first=True, line=1.0)
    tf = tb(slide, MX + CW - 1.2, 7.02, 1.2, 0.3)
    para(tf, "%d / %d" % (idx, total), F_FOOT, rgb("faint"),
         align=PP_ALIGN.RIGHT, first=True, line=1.0)


def header(slide, eyebrow, title):
    y = TOP
    if eyebrow:
        tf = tb(slide, MX, y, CW, 0.3)
        para(tf, eyebrow, F_EYEBROW, rgb("muted"), bold=True, first=True, line=1.0)
        y += 0.38
    if title:
        tf = tb(slide, MX, y, CW, 0.8)
        para(tf, title, F_TITLE, rgb("ink"), bold=True, first=True, line=1.08)
        y += 0.92
    return y


# --------------------------------------------------------------------------
# 各版式
# --------------------------------------------------------------------------
def _set_strike(run):
    run._r.get_or_add_rPr().set("strike", "sngStrike")


def lv_cover(slide, d, ctx):
    tf = tb(slide, MX, 2.05, CW, 0.4)
    para(tf, d.get("eyebrow", ""), 15, rgb("muted"), bold=True,
         first=True, line=1.0)
    tf = tb(slide, MX, 2.5, CW, 1.6)
    para(tf, d["title"], 44, rgb("ink"), bold=True, first=True, line=1.04)
    if d.get("subtitle"):
        tf = tb(slide, MX, 3.6, CW * 0.80, 0.9)
        para(tf, d["subtitle"], 20, rgb("muted"), first=True, line=1.3)
    # 封面也要有一个视觉：一个「改前 → 改后」的大对比，唯一的琥珀落在改后的数上
    if d.get("hero_to"):
        tf = tb(slide, MX, 4.62, CW, 1.0)
        p = tf.paragraphs[0]
        p.line_spacing = 1.0
        if d.get("hero_from"):
            _set_strike(_run(p, d["hero_from"], 30, rgb("muted"), True))
            _run(p, "  \u2192  ", 24, rgb("faint"))
        _run(p, d["hero_to"], 56, rgb("accent_ink"), True)
        if d.get("hero_cap"):
            tf2 = tb(slide, MX, 5.52, CW, 0.3)
            para(tf2, d["hero_cap"], 13, rgb("faint"), first=True, line=1.0)
    hline(slide, MX, 6.18, CW, rgb("line"), 1.0)
    tf = tb(slide, MX, 6.34, CW, 0.4)
    para(tf, d.get("meta", ""), 14, rgb("faint"), first=True, line=1.0)


def lv_section(slide, d, ctx):
    tf = tb(slide, MX, 2.85, CW, 0.4)
    para(tf, d.get("eyebrow", ""), F_EYEBROW, rgb("accent_ink"), bold=True,
         first=True, line=1.0)
    tf = tb(slide, MX, 3.25, CW, 1.1)
    para(tf, d["title"], 38, rgb("ink"), bold=True, first=True, line=1.05)
    if d.get("subtitle"):
        tf = tb(slide, MX, 4.55, CW * 0.78, 0.9)
        para(tf, d["subtitle"], F_BODY, rgb("muted"), first=True, line=1.4)


def lv_statement(slide, d, ctx):
    if d.get("eyebrow"):
        tf = tb(slide, MX, 1.95, CW, 0.4)
        para(tf, d["eyebrow"], F_EYEBROW, rgb("muted"), bold=True,
             first=True, line=1.0)
    text, acc = d["text"], d.get("accent")
    tf = tb(slide, MX, 2.55, CW * 0.94, 2.4)
    p = tf.paragraphs[0]
    p.line_spacing = 1.15
    if acc and acc in text:
        pre, _, post = text.partition(acc)
        if pre:
            _run(p, pre, 33, rgb("ink"), True)
        _run(p, acc, 33, rgb("accent_ink"), True)
        if post:
            _run(p, post, 33, rgb("ink"), True)
    else:
        _run(p, text, 33, rgb("ink"), True)
    if d.get("note"):
        tf = tb(slide, MX, 5.05, CW * 0.86, 0.9)
        para(tf, d["note"], 16, rgb("muted"), first=True, line=1.45)


def lv_bullets(slide, d, ctx):
    y = header(slide, d.get("eyebrow"), d.get("title"))
    items = d.get("items", [])
    aside = d.get("aside")
    left_w = CW * 0.60 if aside else CW
    rowh = 0.62
    for i, it in enumerate(items):
        it = it if isinstance(it, dict) else {"text": it}
        yy = y + 0.24 + i * rowh
        oval(slide, MX + 0.09, yy + 0.16, 0.12, rgb("faint"))
        tf = tb(slide, MX + 0.36, yy, left_w - 0.46, 0.56)
        p = tf.paragraphs[0]
        p.line_spacing = 1.08
        _run(p, it["text"], 19, rgb("ink"))
        if it.get("meta"):
            para(tf, it["meta"], 14, rgb("muted"), line=1.1, space_before=2)
    if aside:
        ax, aw = MX + CW * 0.64, CW * 0.36
        note = aside.get("note") or ""
        cpl = max(8, int((aw - 0.68) / 0.20))
        nl = int(math.ceil(len(note) / float(cpl))) if note else 0
        card_h = 1.62 + 0.28 * max(1, nl) + 0.28   # 卡片贴着内容长，别拉满整页
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, ax, y + 0.15, aw, card_h,
              fill=rgb("surface"), line=rgb("line"), radius=0.05)
        tf = tb(slide, ax + 0.34, y + 0.56, aw - 0.68, 0.4)
        para(tf, aside.get("label", ""), F_EYEBROW, rgb("muted"), bold=True,
             first=True, line=1.0)
        tf = tb(slide, ax + 0.34, y + 0.92, aw - 0.68, 0.9)
        para(tf, aside["value"], 40, rgb("accent_ink"), bold=True,
             first=True, line=1.0)
        if note:
            tf = tb(slide, ax + 0.34, y + 1.62, aw - 0.68, card_h - 1.7)
            para(tf, note, 14, rgb("muted"), first=True, line=1.4)


def lv_two_col(slide, d, ctx):
    y = header(slide, d.get("eyebrow"), d.get("title"))
    cols = d.get("cols", [])[:2]
    gap = 0.42
    w = (CW - gap) / 2.0
    cpl = max(8, int((w - 0.64) / 0.22))
    per = [max(1, sum(max(1, int(math.ceil(len(it) / float(cpl))))
                      for it in col.get("items", []))) for col in cols]
    h = 1.10 + 0.44 * max(per + [1]) + 0.30   # 两卡同高，按更高的一列贴内容
    for j, col in enumerate(cols):
        x = MX + j * (w + gap)
        tone = col.get("tone", "note")
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y + 0.10, w, h,
              fill=rgb(TONE_WASH.get(tone, C["fill"])), line=None, radius=0.045)
        tf = tb(slide, x + 0.32, y + 0.44, w - 0.64, 0.5)
        para(tf, col["title"], 19, rgb(TONE_INK.get(tone, C["muted"])),
             bold=True, first=True, line=1.1)
        tf = tb(slide, x + 0.32, y + 1.10, w - 0.64, h - 1.35)
        for k, line in enumerate(col.get("items", [])):
            para(tf, line, 16, rgb("text"), first=(k == 0),
                 line=1.45, space_after=8)


def _patch_txpr(root, size, color):
    """给图表里所有文字补上 a:ea（CJK），并统一字号/颜色。"""
    for defRPr in root.iter(qn("a:defRPr")):
        if size is not None:
            defRPr.set("sz", str(int(size * 100)))
        _set_face(defRPr, "a:latin", LATIN)
        _set_face(defRPr, "a:ea", EA)
        _set_face(defRPr, "a:cs", LATIN)
        if color is not None:
            for e in defRPr.findall(qn("a:solidFill")):
                defRPr.remove(e)
            sf = defRPr.makeelement(qn("a:solidFill"), {})
            clr = sf.makeelement(qn("a:srgbClr"), {"val": color})
            sf.append(clr)
            # solidFill 必须排在 latin 之前
            latin = defRPr.find(qn("a:latin"))
            if latin is not None:
                latin.addprevious(sf)
            else:
                defRPr.append(sf)


def lv_chart(slide, d, ctx):
    y = header(slide, d.get("eyebrow"), d.get("title"))
    spec = d["chart"]
    cd = CategoryChartData()
    cd.categories = spec["categories"]
    for s in spec["series"]:
        cd.add_series(s["name"], tuple(s["values"]))
    ctype = CHARTS.get(spec.get("type", "column"), XL_CHART_TYPE.COLUMN_CLUSTERED)
    ch_w = CW * 0.63
    gf = slide.shapes.add_chart(ctype, Inches(MX), Inches(y + 0.12),
                                Inches(ch_w), Inches(4.15), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    plot = ch.plots[0]
    plot.gap_width = 70
    try:
        plot.overlap = -12
    except Exception:
        pass
    # 两条以内不靠彩度区分：改前浅灰、改后深灰，去色也分得开（同 HTML 侧）
    series_colors = [C["faint"], C["text"], C["chart3"], C["chart4"]]
    for i, s in enumerate(ch.series):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(series_colors[i % len(series_colors)])
    # 读数差两三个数量级（十几秒 vs 一秒），改后那根柱会贴着零点——不能靠
    # 对数轴（LibreOffice 会把数据标签压在轴的基线处，正好糊掉最该读的那个数），
    # 改成在柱顶直标数值，「14.5 → 0.91」照样读得出来。
    try:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_value = True
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(11)
        dl.font.name = LATIN
        dl.font.color.rgb = rgb("text")
    except Exception:
        pass
    ch.font.size = Pt(12)
    ch.font.name = LATIN
    ch.legend.font.size = Pt(12)
    ch.legend.font.name = LATIN
    for ax in (ch.category_axis, ch.value_axis):
        ax.tick_labels.font.size = Pt(12)
        ax.tick_labels.font.name = LATIN
        ax.has_major_gridlines = ax is ch.value_axis
        if ax is ch.value_axis:
            try:
                ax.major_gridlines.format.line.color.rgb = rgb("line")
                ax.major_gridlines.format.line.width = Pt(0.5)
            except Exception:
                pass
        ax.format.line.color.rgb = rgb("line2")
    _patch_txpr(ch._chartSpace, 12, C["muted"])
    # 右侧结论卡：贴着正文高度收，不留一大片空
    tx, tw = MX + CW * 0.67, CW * 0.33
    takeaway = d.get("takeaway", "")
    cpl = max(8, int((tw - 0.60) / 0.22))
    tl = max(1, int(math.ceil(len(takeaway) / float(cpl)))) if takeaway else 0
    card_h = 0.56 + 0.34 * max(1, tl) + 0.28
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, tx, y + 0.25, tw, card_h,
          fill=rgb("fill"), line=rgb("line"), radius=0.05)
    tf = tb(slide, tx + 0.30, y + 0.56, tw - 0.60, card_h - 0.5)
    para(tf, takeaway, 16, rgb("text"), first=True, line=1.5)
    if d.get("caption"):
        tf = tb(slide, MX, 6.42, CW, 0.45)
        para(tf, d["caption"], F_CAPTION, rgb("muted"), first=True, line=1.3)


def _cell_border(cell, edge, color_hex, w_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"L": "a:lnL", "R": "a:lnR", "T": "a:lnT", "B": "a:lnB"}[edge]
    for e in tcPr.findall(qn(tag)):
        tcPr.remove(e)
    ln = tcPr.makeelement(qn(tag), {"w": str(int(w_pt * 12700)),
                                    "cap": "flat", "cmpd": "sng", "algn": "ctr"})
    sf = ln.makeelement(qn("a:solidFill"), {})
    clr = sf.makeelement(qn("a:srgbClr"), {"val": color_hex})
    sf.append(clr)
    ln.append(sf)
    fill = None
    for t in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill",
              "a:pattFill", "a:grpFill"):
        el = tcPr.find(qn(t))
        if el is not None:
            fill = el
            break
    if fill is not None:
        fill.addprevious(ln)
    else:
        tcPr.append(ln)


def lv_table(slide, d, ctx):
    y = header(slide, d.get("eyebrow"), d.get("title"))
    td = d["table"]
    cols, rows = td["columns"], td["rows"]
    nrow, ncol = len(rows) + 1, len(cols)
    tw = CW
    th = min(0.52 + 0.56 * len(rows), 4.35)
    gf = slide.shapes.add_table(nrow, ncol, Inches(MX), Inches(y + 0.14),
                                Inches(tw), Inches(th))
    table = gf.table
    # 关掉默认的主题表格样式，表格线自己画
    tblPr = table._tbl.find(qn("a:tblPr"))
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = tblPr.makeelement(qn("a:tableStyleId"), {})
        tblPr.append(sid)
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # No Style, No Grid

    # 列宽：首列更宽，其余等分
    first_w = tw * 0.34
    rest_w = (tw - first_w) / max(1, ncol - 1)
    for j in range(ncol):
        table.columns[j].width = Inches(first_w if j == 0 else rest_w)
    for i in range(nrow):
        table.rows[i].height = Inches(th / nrow)

    aligns = td.get("col_align", ["l"] + ["r"] * (ncol - 1))
    amap = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}

    for j, name in enumerate(cols):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb("fill")
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.14)
        cell.margin_right = Inches(0.14)
        tf = cell.text_frame
        tf.word_wrap = True
        para(tf, name, F_EYEBROW, rgb("muted"), bold=True, first=True,
             align=amap.get(aligns[j] if j < len(aligns) else "l", PP_ALIGN.LEFT),
             line=1.0)
        _cell_border(cell, "B", C["line2"], 0.75)

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb("surface")
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.14)
            cell.margin_right = Inches(0.14)
            tf = cell.text_frame
            tf.word_wrap = True
            align = amap.get(aligns[j] if j < len(aligns) else "l", PP_ALIGN.LEFT)
            p = tf.paragraphs[0]
            p.alignment = align
            p.line_spacing = 1.0
            if isinstance(val, dict):
                tone = val.get("tone", "neutral")
                _run(p, "●  ", 12, rgb(TONE_MARK.get(tone, C["faint"])))
                _run(p, val.get("text", ""), 14, rgb(TONE_INK.get(tone, C["text"])), True)
            else:
                _run(p, str(val), 14, rgb("text"), j == 0)
            _cell_border(cell, "B", C["line"], 0.75)

    if d.get("caption"):
        tf = tb(slide, MX, y + 0.25 + th, CW, 0.5)
        para(tf, d["caption"], F_CAPTION, rgb("muted"), first=True, line=1.3)


def lv_timeline(slide, d, ctx):
    y = header(slide, d.get("eyebrow"), d.get("title"))
    steps = d.get("steps", [])
    n = len(steps)
    ly = 3.7
    # 两端各留出标签盒的一半，话题编号不会顶到版心外
    x0, x1 = MX + 1.85, MX + CW - 1.85
    span = x1 - x0
    hline(slide, x0, ly, span, rgb("line2"), 2.0)
    for i, st in enumerate(steps):
        cx = x0 + span * i / (n - 1) if n > 1 else x0 + span / 2.0
        col = "accent" if i == 0 else "ink"
        oval(slide, cx, ly, 0.26, rgb(col))
        tf = tb(slide, cx - 1.75, ly - 1.7, 3.5, 1.3)
        para(tf, st["label"], 20, rgb("ink"), bold=True,
             align=PP_ALIGN.CENTER, first=True, line=1.15)
        if st.get("desc"):
            para(tf, st["desc"], 14, rgb("muted"), align=PP_ALIGN.CENTER,
                 line=1.3, space_before=4)
    if d.get("axis_note"):
        tf = tb(slide, MX, ly + 1.5, CW, 0.5)
        para(tf, d["axis_note"], F_CAPTION, rgb("faint"),
             align=PP_ALIGN.CENTER, first=True, line=1.3)


def lv_closing(slide, d, ctx):
    tf = tb(slide, MX, 2.55, CW, 0.4)
    para(tf, d.get("eyebrow", ""), 15, rgb("muted"), bold=True,
         first=True, line=1.0)
    tf = tb(slide, MX, 3.0, CW, 1.1)
    para(tf, d["title"], 40, rgb("ink"), bold=True, first=True, line=1.05)
    if d.get("lines"):
        tf = tb(slide, MX, 4.35, CW * 0.82, 1.6)
        for i, ln in enumerate(d["lines"]):
            # 一条行可以写成 {"text": "...", "accent": true}，
            # 全页唯一的琥珀落到「下一步 / 拍板」那一行，不放在眉标上
            if isinstance(ln, dict):
                text = ln.get("text", "")
                acc = bool(ln.get("accent"))
            else:
                text, acc = ln, False
            para(tf, text, F_BODY,
                 rgb("accent_ink") if acc else rgb("muted"),
                 bold=acc, first=(i == 0), line=1.5, space_after=4)
    hline(slide, MX, 6.18, CW, rgb("line"), 1.0)


LAYOUTS = {
    "cover": lv_cover, "section": lv_section, "statement": lv_statement,
    "bullets": lv_bullets, "two-col": lv_two_col, "chart": lv_chart,
    "table": lv_table, "timeline": lv_timeline, "closing": lv_closing,
}
# 有底栏的版式（封面 / 结尾留白，不加页码）
FOOTERED = {"statement", "bullets", "two-col", "chart", "table",
            "timeline", "section"}


def build(spec, out_path):
    prs = new_deck()
    slides = spec.get("slides", [])
    total = len(slides)
    name = spec.get("title", "")
    for idx, d in enumerate(slides, start=1):
        layout = d.get("layout", "bullets")
        slide = add_slide(prs)
        LAYOUTS.get(layout, lv_bullets)(slide, d, {"idx": idx, "total": total})
        if layout in FOOTERED:
            footer(slide, idx, total, name)
        set_notes(slide, d.get("notes"))
    prs.save(out_path)
    return total


def main():
    if len(sys.argv) < 3:
        sys.stderr.write("用法: python build_deck.py spec.json out.pptx\n")
        return 2
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    n = build(spec, sys.argv[2])
    print("built %d slides -> %s" % (n, sys.argv[2]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
