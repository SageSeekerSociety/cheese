"""Read office files into text that says where every piece came from.

    uv run --with pdfplumber --with python-docx --with python-pptx --with openpyxl \
        python3 read.py 材料.pdf 方案.docx 汇报.pptx 预算.xlsx \
        [--pages 3-7] [--media 目录]

Each file comes back as markdown with its positions spelled out — PDF pages,
Word paragraph numbers and heading path, table cells, slide numbers, sheet
cells with the formula and the computed value kept apart — so a sentence in a
report can cite exactly where it came from. Every file ends with what was NOT
read: a scanned page, a chart, SmartArt, a formula with no computed value, a
part cut off by the size limit. Nothing is left out silently.

Exit code: 0 when every file was opened, 1 when any could not be opened at all.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
CHART_URI = "http://schemas.openxmlformats.org/drawingml/2006/chart"
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"


class Out:
    def __init__(self, limit: int):
        self.lines: list[str] = []
        self.missing: list[str] = []
        self.limit = limit
        self.size = 0
        self.cut_at: str | None = None

    def add(self, text: str, where: str = "") -> bool:
        """False once the size limit is reached; the caller stops there."""
        if self.cut_at is not None:
            return False
        if self.size + len(text) > self.limit:
            self.cut_at = where or "这里"
            return False
        self.lines.append(text)
        self.size += len(text) + 1
        return True

    def miss(self, text: str) -> None:
        self.missing.append(text)


def cell_text(value) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\n", " ").strip()


def table_md(rows: list[list], caption: str) -> str:
    if not rows:
        return f"{caption}（空表）"
    width = max(len(r) for r in rows)
    head = ["行"] + [f"C{i + 1}" for i in range(width)]
    lines = [caption, "", "| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
    for n, row in enumerate(rows, 1):
        cells = [cell_text(c) for c in row] + [""] * (width - len(row))
        lines.append(f"| R{n} | " + " | ".join(cells) + " |")
    return "\n".join(lines)


def page_range(spec: str | None, total: int) -> list[int]:
    if not spec:
        return list(range(1, total + 1))
    pages: set[int] = set()
    for part in spec.split(","):
        if "-" in part:
            a, b = part.split("-", 1)
            pages.update(range(int(a), int(b) + 1))
        elif part.strip():
            pages.add(int(part))
    return [p for p in sorted(pages) if 1 <= p <= total]


# ── PDF ─────────────────────────────────────────────────────────────────────


def read_pdf(path: Path, out: Out, args) -> None:
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        total = len(pdf.pages)
        wanted = page_range(args.pages, total)
        out.add(f"# 《{path.name}》（PDF，共 {total} 页）")
        if len(wanted) < total:
            out.miss(f"只读了第 {args.pages} 页（共 {total} 页），其余页没有读")
        for number in wanted:
            page = pdf.pages[number - 1]
            text = (page.extract_text() or "").strip()
            images = len(page.images)
            if not out.add(f"\n## 第 {number} 页\n", f"第 {number} 页"):
                break
            if len(text) < 20 and images:
                out.miss(
                    f"第 {number} 页：几乎没有文字层、有 {images} 张图片，"
                    "可能是扫描页，里面的内容没有读到"
                )
            elif not text:
                out.miss(f"第 {number} 页：没有抽出任何文字")
            if text and not out.add(text, f"第 {number} 页"):
                break
            for i, table in enumerate(page.extract_tables() or [], 1):
                if not out.add(
                    "\n" + table_md(table, f"表 {number}-{i}（第 {number} 页）"),
                    f"第 {number} 页的表 {i}",
                ):
                    break
            if images and len(text) >= 20:
                out.add(f"（本页另有 {images} 张图片，图片里的内容没有识别）")
                out.miss(f"第 {number} 页的 {images} 张图片：图里的文字和数据没有识别")


# ── Word ────────────────────────────────────────────────────────────────────


def _alt_texts(element) -> list[str]:
    found = []
    for doc_pr in element.iter(
        "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
    ):
        found.append(doc_pr.get("descr") or doc_pr.get("name") or "")
    return found


def read_docx(path: Path, out: Out, args) -> None:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    body = document.element.body
    out.add(f"# 《{path.name}》（Word）")
    heading: list[str] = []
    paragraph_no = 0
    table_no = 0
    image_no = 0
    media = Path(args.media) if args.media else None
    rels = document.part.rels

    for child in body.iterchildren():
        tag = child.tag
        if tag == f"{W}p":
            paragraph_no += 1
            p = Paragraph(child, document)
            style = (p.style.name if p.style is not None else "") or ""
            text = p.text.strip()
            level = None
            match = re.match(r"(?:Heading|标题)\s*(\d)", style)
            if match:
                level = int(match.group(1))
            elif style == "Title":
                level = 0
            if level is not None and text:
                heading = heading[: max(level - 1, 0)] + [text]
                ok = out.add(f"\n{'#' * min(level + 2, 6)} {text}（¶{paragraph_no}）")
            elif text:
                is_list = child.find(f"{W}pPr/{W}numPr") is not None or "List" in style
                prefix = "- " if is_list else ""
                ok = out.add(f"¶{paragraph_no} {prefix}{text}", f"¶{paragraph_no}")
            else:
                ok = True
            for alt in _alt_texts(child):
                image_no += 1
                note = (
                    f"（¶{paragraph_no} 有图片 {image_no}"
                    + (f"：{alt}" if alt else "")
                    + "）"
                )
                out.add(note)
            for blip in child.iter(f"{A}blip"):
                rid = blip.get(f"{R}embed")
                if media is not None and rid in rels:
                    media.mkdir(parents=True, exist_ok=True)
                    target = rels[rid].target_part
                    name = f"{path.stem}-¶{paragraph_no}-{Path(target.partname).name}"
                    (media / name).write_bytes(target.blob)
            for frame in child.iter(f"{A}graphicData"):
                uri = frame.get("uri")
                if uri == CHART_URI:
                    out.miss(f"¶{paragraph_no} 附近有一个图表：图表数据没有读")
                elif uri == DIAGRAM_URI:
                    out.miss(f"¶{paragraph_no} 附近有 SmartArt：里面的文字没有读")
            if child.find(f".//{W}object") is not None:
                out.miss(f"¶{paragraph_no} 有嵌入对象（OLE）：没有读")
            boxes = [
                "".join(t.text or "" for t in box.iter(f"{W}t"))
                for box in child.iter(f"{W}txbxContent")
            ]
            for box in filter(None, boxes):
                out.add(f"¶{paragraph_no} [文本框] {box}")
            if not ok:
                break
        elif tag == f"{W}tbl":
            table_no += 1
            table = Table(child, document)
            rows = [[c.text for c in row.cells] for row in table.rows]
            where = f"表 {table_no}（在 ¶{paragraph_no} 之后"
            where += f"，所在章节：{' > '.join(heading)}）" if heading else "）"
            if not out.add("\n" + table_md(rows, where), f"表 {table_no}"):
                break
    if image_no:
        extracted = (
            f"已导出到 {media}，可以打开看" if media else "加 --media 目录 可以导出来看"
        )
        out.miss(f"{image_no} 张图片：图里的内容没有识别（{extracted}）")
    for section_no, section in enumerate(document.sections, 1):
        for kind, part in (("页眉", section.header), ("页脚", section.footer)):
            text = " ".join(p.text for p in part.paragraphs if p.text.strip())
            if text:
                out.add(f"[第 {section_no} 节{kind}] {text}")
    if any("comments" in rel.reltype for rel in rels.values()):
        out.miss("文档里有批注：批注内容没有读")
    if any("footnotes" in rel.reltype for rel in rels.values()):
        out.miss("文档里有脚注：脚注内容没有读")


# ── PowerPoint ──────────────────────────────────────────────────────────────


def _shapes(shapes):
    for shape in shapes:
        if shape.shape_type == 6 and hasattr(shape, "shapes"):  # group
            yield from _shapes(shape.shapes)
        else:
            yield shape


def read_pptx(path: Path, out: Out, args) -> None:
    from pptx import Presentation

    deck = Presentation(str(path))
    out.add(f"# 《{path.name}》（演示文稿，共 {len(deck.slides)} 页）")
    media = Path(args.media) if args.media else None
    wanted = set(page_range(args.pages, len(deck.slides)))
    if len(wanted) < len(deck.slides):
        out.miss(f"只读了第 {args.pages} 页（共 {len(deck.slides)} 页）")
    for number, slide in enumerate(deck.slides, 1):
        if number not in wanted:
            continue
        title = (
            slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        )
        if not out.add(
            f"\n## 幻灯片 {number}" + (f"：{title}" if title else ""),
            f"幻灯片 {number}",
        ):
            break
        for shape in _shapes(slide.shapes):
            name = shape.name
            if shape.has_text_frame and shape.text_frame.text.strip():
                if shape == slide.shapes.title:
                    continue
                text = "\n".join(
                    ("  " * p.level + "- " if p.level else "")
                    + "".join(r.text for r in p.runs)
                    for p in shape.text_frame.paragraphs
                    if "".join(r.text for r in p.runs).strip()
                )
                out.add(f"[{name}] {text}")
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text for c in row.cells] for row in shape.table.rows]
                out.add("\n" + table_md(rows, f"表格「{name}」（幻灯片 {number}）"))
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                chart = shape.chart
                lines = [f"[图表「{name}」，幻灯片 {number}]"]
                try:
                    categories = list(chart.plots[0].categories)
                    lines.append("类别：" + "、".join(map(str, categories)))
                    for series in chart.series:
                        values = ", ".join(str(v) for v in series.values)
                        lines.append(f"系列「{series.name}」：{values}")
                except Exception as exc:  # noqa: BLE001 — reported, not hidden
                    out.miss(f"幻灯片 {number} 的图表「{name}」：数据没读出来（{exc}）")
                out.add("\n".join(lines))
            elif shape.shape_type == 13:  # picture
                alt = shape._element.xpath("./p:nvPicPr/p:cNvPr/@descr")
                out.add(
                    f"[图片「{name}」" + (f"：{alt[0]}" if alt and alt[0] else "") + "]"
                )
                out.miss(f"幻灯片 {number} 的图片「{name}」：图里的内容没有识别")
                if media is not None:
                    media.mkdir(parents=True, exist_ok=True)
                    image = shape.image
                    (
                        media / f"{path.stem}-幻灯片{number}-{name}.{image.ext}"
                    ).write_bytes(image.blob)
            elif shape.shape_type == 7 or DIAGRAM_URI in shape._element.xml:
                out.miss(f"幻灯片 {number} 的「{name}」是 SmartArt/嵌入对象：没有读")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.add(f"[备注] {notes}")


# ── Excel ───────────────────────────────────────────────────────────────────


def read_xlsx(path: Path, out: Out, args) -> None:
    from openpyxl import load_workbook

    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    out.add(f"# 《{path.name}》（表格，{len(formulas.sheetnames)} 个工作表）")
    uncomputed: list[str] = []
    for name in formulas.sheetnames:
        ws, wv = formulas[name], values[name]
        state = "" if ws.sheet_state == "visible" else f"，{ws.sheet_state}"
        if not out.add(
            f"\n## 工作表「{name}」（范围 {ws.dimensions}{state}）", f"工作表「{name}」"
        ):
            break
        if ws.merged_cells.ranges:
            out.add("合并单元格：" + "、".join(str(r) for r in ws.merged_cells.ranges))
        shown = 0
        total = 0
        stop = False
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                total += 1
                if shown >= args.max_cells:
                    continue
                ref = f"{name}!{cell.coordinate}"
                raw = cell.value
                if isinstance(raw, str) and raw.startswith("="):
                    cached = wv[cell.coordinate].value
                    if cached is None:
                        uncomputed.append(ref)
                        line = f"{cell.coordinate} = （没有计算结果） ← 公式 {raw}"
                    else:
                        line = f"{cell.coordinate} = {cached} ← 公式 {raw}"
                elif hasattr(raw, "text"):  # array formula
                    computed = wv[cell.coordinate].value
                    line = f"{cell.coordinate} = {computed} ← 数组公式 {raw.text}"
                else:
                    line = f"{cell.coordinate} = {raw}"
                shown += 1
                if not out.add(line, ref):
                    stop = True
                    break
            if stop:
                break
        if stop:
            break
        if total > shown:
            out.miss(
                f"工作表「{name}」有 {total} 个非空单元格，只列了前 {shown} 个"
                "（加 --max-cells 调大，或用 openpyxl 按区域读）"
            )
        for chart in getattr(ws, "_charts", []):
            title = ""
            try:
                title = "".join(r.t for p in chart.title.tx.rich.p for r in (p.r or []))
            except AttributeError:
                pass
            refs = []
            for series in chart.series:
                if series.val is not None and series.val.numRef is not None:
                    refs.append(series.val.numRef.f)
            out.add(
                f"[图表 {type(chart).__name__}{'「' + title + '」' if title else ''}，"
                f"数据 {', '.join(refs) or '未知'}]"
            )
        if getattr(ws, "_pivots", None):
            out.miss(f"工作表「{name}」有数据透视表：透视表本身没有读")
    if uncomputed:
        shown = "、".join(uncomputed[:10]) + ("等" if len(uncomputed) > 10 else "")
        out.miss(
            f"{len(uncomputed)} 个公式没有计算结果（{shown}）："
            "先 `cheese recalc` 重算再读，不要把公式当成数字"
        )


READERS = {
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".pptx": read_pptx,
    ".xlsx": read_xlsx,
    ".xlsm": read_xlsx,
}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("files", nargs="+")
    parser.add_argument("--pages", help="只读这些页，例如 3-7,10（PDF 和演示文稿）")
    parser.add_argument(
        "--media", help="把 Word/演示文稿里的图片导出到这个目录，好打开看"
    )
    parser.add_argument(
        "--max-chars", type=int, default=60000, help="每份文件最多输出多少字"
    )
    parser.add_argument(
        "--max-cells", type=int, default=2000, help="每个工作表最多列多少格"
    )
    args = parser.parse_args(argv)

    failed = False
    summary = []
    for name in args.files:
        path = Path(name)
        out = Out(args.max_chars)
        reader = READERS.get(path.suffix.lower())
        if not path.is_file():
            summary.append(f"- 《{path.name}》：找不到这个文件")
            failed = True
            continue
        if reader is None:
            hint = (
                "老格式先 `cheese convert` 升级"
                if path.suffix.lower() in (".doc", ".ppt", ".xls")
                else "这个脚本不读这种格式"
            )
            summary.append(f"- 《{path.name}》：没有读（{hint}）")
            failed = True
            continue
        try:
            reader(path, out, args)
        except Exception as exc:  # noqa: BLE001 — the reader of a corrupt file must say so
            summary.append(f"- 《{path.name}》：打不开（{type(exc).__name__}: {exc}）")
            failed = True
            continue
        if out.cut_at is not None:
            out.miss(
                f"输出到 {out.cut_at} 就停了（超过 {args.max_chars} 字）；"
                "后面的部分没有读"
            )
        print("\n".join(out.lines))
        print(f"\n## 《{path.name}》没有读到的部分\n")
        print(
            "\n".join(f"- {m}" for m in out.missing)
            if out.missing
            else "- 无：这份文件的文字、表格都已列出"
        )
        print()
        summary.append(
            f"- 《{path.name}》：已读"
            + (f"，{len(out.missing)} 处没读到（见上）" if out.missing else "")
        )
    print("# 读取结果\n")
    print("\n".join(summary))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
