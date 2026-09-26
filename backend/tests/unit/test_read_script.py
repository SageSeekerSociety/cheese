"""`read.py` says where each piece came from and what it did not read.

A report that cites 「第 3 页」 or 「汇总!B2」 is only honest if the reader put that
position there, and a summary is only honest if a scanned page, a picture or an
uncomputed formula was reported instead of skipped. Run the way the agent runs
it: `uv run --with …` on real files.
"""

from __future__ import annotations

import shutil
import subprocess
import textwrap

import pytest

from app.domain.agent.skills import _NATIVE_SKILL_SRC

READ = _NATIVE_SKILL_SRC / "documents" / "scripts" / "read.py"
WITH = ["--with", "pdfplumber", "--with", "python-docx", "--with", "python-pptx"]
WITH += ["--with", "openpyxl", "--with", "fpdf2"]

pytestmark = pytest.mark.skipif(shutil.which("uv") is None, reason="needs uv")

MAKE = textwrap.dedent(
    """
    import struct, zlib
    from docx import Document
    from docx.shared import Inches
    from fpdf import FPDF
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches as PI

    def png():
        raw = b"".join(b"\\x00" + b"\\xc8\\x1e\\x1e" * 40 for _ in range(20))
        def chunk(t, d):
            crc = struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
            return struct.pack(">I", len(d)) + t + d + crc
        head = struct.pack(">IIBBBBB", 40, 20, 8, 2, 0, 0, 0)
        return (b"\\x89PNG\\r\\n\\x1a\\n" + chunk(b"IHDR", head)
                + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))
    open("logo.png", "wb").write(png())

    pdf = FPDF()
    pdf.add_page(); pdf.set_font("helvetica", size=12)
    pdf.cell(0, 10, "Budget used: 33 (external review)")
    pdf.add_page(); pdf.image("logo.png", w=180)
    pdf.output("review.pdf")

    d = Document()
    d.add_heading("Background", 1)
    d.add_paragraph("Twelve tasks were finished.")
    d.add_heading("Budget", 2)
    d.add_paragraph("Budget used: 31.5")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "module"; t.cell(0, 1).text = "progress"
    t.cell(1, 0).text = "analysis"; t.cell(1, 1).text = "45%"
    d.add_picture("logo.png", width=Inches(1))
    d.save("report.docx")

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = "Trend"
    cd = CategoryChartData(); cd.categories = ["Jul", "Aug"]
    cd.add_series("done", (3, 4))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, PI(1), PI(1.5), PI(6), PI(4), cd)
    s.notes_slide.notes_text_frame.text = "stress the budget"
    p.save("deck.pptx")

    wb = Workbook(); a = wb.active; a.title = "input"
    a.append(["month", "cost"]); a.append(["Jul", 10.5]); a.append(["Aug", 11])
    b = wb.create_sheet("sum"); b["B2"] = "=SUM(input!B2:B3)"
    wb.save("budget.xlsx")
    """
)


@pytest.fixture(scope="module")
def samples(tmp_path_factory):
    where = tmp_path_factory.mktemp("samples")
    subprocess.run(
        ["uv", "run", "-q", *WITH, "python3", "-c", MAKE], cwd=where, check=True
    )
    (where / "broken.pdf").write_text("not a pdf")
    return where


def _read(samples, *files):
    result = subprocess.run(
        ["uv", "run", "-q", *WITH, "python3", str(READ), *files],
        cwd=samples,
        capture_output=True,
        text=True,
    )
    return result.returncode, result.stdout


def test_every_piece_carries_its_position(samples):
    code, out = _read(samples, "review.pdf", "report.docx", "deck.pptx", "budget.xlsx")
    assert code == 0, out
    assert "## 第 1 页" in out and "Budget used: 33" in out
    assert "¶4 Budget used: 31.5" in out or "Budget used: 31.5" in out
    assert "所在章节：Background > Budget" in out
    assert "| R2 | analysis | 45% |" in out
    assert "## 幻灯片 1：Trend" in out
    assert "系列「done」：3.0, 4.0" in out
    assert "[备注] stress the budget" in out


def test_a_formula_without_a_value_is_not_passed_off_as_a_number(samples):
    _code, out = _read(samples, "budget.xlsx")
    assert "B2 = （没有计算结果） ← 公式 =SUM(input!B2:B3)" in out
    assert "sum!B2" in out.split("没有读到的部分")[-1]


def test_what_was_not_read_is_listed(samples):
    _code, out = _read(samples, "review.pdf", "report.docx")
    unread = out.split("# 《report.docx》")[0].split("没有读到的部分")[-1]
    assert "第 2 页" in unread and "扫描" in unread
    assert "图片" in out.split("# 《report.docx》")[1].split("没有读到的部分")[-1]


def test_a_file_that_cannot_be_read_is_reported_and_fails_the_run(samples):
    code, out = _read(samples, "budget.xlsx", "broken.pdf", "missing.docx", "old.doc")
    assert code == 1
    summary = out.split("# 读取结果")[-1]
    assert "《broken.pdf》：打不开" in summary
    assert "《missing.docx》：找不到" in summary
    assert "《budget.xlsx》：已读" in summary
